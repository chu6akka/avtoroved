import json
from hashlib import sha256
from zipfile import ZipFile

import pytest
from docx import Document as WordDocument

from authoroved_core.core.case_repository import CaseRepository
from authoroved_core.core.comparison import compare_results
from authoroved_core.core.document import Document
from authoroved_core.core.models import AnalysisResult, Candidate, Metric, ReviewStatus, Span
from authoroved_core.core.report_service import ReportError, ReportService


def source_document(number, text):
    data = text.encode("utf-8")
    return Document(
        id=f"doc-{number}", name=f"Текст {number}.txt", text=text, original_bytes=data,
        file_sha256=sha256(data).hexdigest(), text_sha256=sha256(data).hexdigest(),
        encoding="utf-8", import_note="Пробелы и переносы сохранены.",
    )


def result(document, accepted, rejected, new):
    candidates = []
    for index, (fragment, status) in enumerate([
        (accepted, ReviewStatus.ACCEPTED),
        (rejected, ReviewStatus.REJECTED),
        (new, ReviewStatus.NEW),
    ]):
        start = document.text.index(fragment)
        candidate = Candidate(
            f"{document.id}-{index}", document.id, "Наблюдение", "Орфография",
            "Автоматическая рекомендация", fragment, Span(start, start + len(fragment)),
            "RULE", status=status,
        )
        candidates.append(candidate)
    return AnalysisResult(
        document_id=document.id,
        candidates=candidates,
        metrics=[
            Metric("Слова", "8", "Слова по разметке Stanza.", "Количественные показатели"),
            Metric("Предложения", "2", "Границы определены Stanza.", "Количественные показатели"),
            Metric("Абзацы", "1", "Непустые строки.", "Структура"),
            Metric("Лексическое разнообразие", "0,75", "Отношение словоформ.", "Лексика"),
            Metric("Вспомогательные глаголы", "1 · 12,5 %", "Код AUX.", "Морфология"),
            Metric(
                "подлежащее (код Stanza: nsubj)", "2 · 25 %", "Служебный код.",
                "Служебная синтаксическая разметка Stanza",
            ),
        ],
        metadata={
            "program_version": "0.2.0",
            "stanza": {"version": "1.14.0", "torch_version": "2.14.0", "device": "cpu"},
            "languagetool": {"version": "6.6", "mode": "local-cli", "build_date": "2026-01-01"},
        },
    )


@pytest.fixture
def ready_case():
    repository = CaseRepository(memory_cost_kib=8192, time_cost=1, parallelism=1)
    case = repository.create()
    first = source_document(1, "принято1 отклонено1 новое1 остальной текст")
    second = source_document(2, "принято2 отклонено2 новое2 другой текст")
    case.materials = [first, second]
    case.results = [
        result(first, "принято1", "отклонено1", "новое1"),
        result(second, "принято2", "отклонено2", "новое2"),
    ]
    return repository, case


def test_docx_contains_auditable_sections_and_only_accepted_observations(tmp_path, ready_case):
    repository, case = ready_case
    target = ReportService(repository).export_docx(tmp_path / "Проект.docx", case)

    document = WordDocument(target)
    text = "\n".join(
        [paragraph.text for paragraph in document.paragraphs]
        + [cell.text for table in document.tables for row in table.rows for cell in row.cells]
    )
    for heading in [
        "Объекты исследования", "Условия автоматизированной обработки",
        "Проверка автоматических рекомендаций", "Сопоставление измерений",
        "Принятые наблюдения", "Ограничения", "Журнал и целостность",
    ]:
        assert heading in text
    assert "принято1" in text and "принято2" in text
    assert "отклонено1" not in text and "новое1" not in text
    assert "не содержит вывода о тождестве или различии авторов" in text
    assert case.materials[0].file_sha256 in text
    assert "Вспомогательные глаголы" not in text
    assert "код Stanza: nsubj" not in text
    assert "Глаголы" in text


def test_verification_package_excludes_sources_by_default_and_checks_hashes(tmp_path, ready_case):
    repository, case = ready_case
    service = ReportService(repository)
    target = service.export_verification_package(tmp_path / "Проверка.zip", case)
    service.verify_verification_package(target)

    with ZipFile(target) as archive:
        names = set(archive.namelist())
        manifest = json.loads(archive.read("manifest.json").decode("utf-8"))
        checksums = json.loads(archive.read("checksums.json").decode("utf-8"))["files"]
        assert not any(name.startswith("sources/") for name in names)
        assert "features.json" in names
        assert manifest["source_documents_included"] is False
        assert set(manifest["contents"]) == names
        for name, expected in checksums.items():
            assert sha256(archive.read(name)).hexdigest() == expected


def test_verification_package_includes_sources_only_by_explicit_choice(tmp_path, ready_case):
    repository, case = ready_case
    target = ReportService(repository).export_verification_package(
        tmp_path / "Полный пакет.zip", case, include_sources=True,
    )

    with ZipFile(target) as archive:
        names = archive.namelist()
        assert "sources/1_Текст 1.txt" in names
        assert "sources/2_Текст 2.txt" in names
        assert archive.read("sources/1_extracted.txt") == case.materials[0].text.encode("utf-8")


def test_modified_verification_entry_is_rejected(tmp_path, ready_case):
    repository, case = ready_case
    service = ReportService(repository)
    target = service.export_verification_package(tmp_path / "Проверка.zip", case)
    with ZipFile(target) as archive:
        entries = {name: archive.read(name) for name in archive.namelist()}
    entries["materials.json"] = b"{}"
    with ZipFile(target, "w") as archive:
        for name, content in entries.items():
            archive.writestr(name, content)

    with pytest.raises(ReportError, match="контрольная сумма"):
        service.verify_verification_package(target)


def test_export_requires_two_analyzed_texts(tmp_path):
    repository = CaseRepository(memory_cost_kib=8192, time_cost=1, parallelism=1)
    case = repository.create()
    with pytest.raises(ReportError, match="два текста"):
        ReportService(repository).export_docx(tmp_path / "Проект.docx", case)
