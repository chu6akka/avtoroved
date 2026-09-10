from dataclasses import replace
from pathlib import Path
from types import SimpleNamespace as NS

import pytest

from authoroved_core.core.document import EncodingChoiceRequired, load_document, sha256_bytes
from authoroved_core.core.models import Candidate, ReviewStatus, Span, Token
from authoroved_core.metrics.basic import calculate_metrics, structural_metrics
from authoroved_core.nlp.languagetool_adapter import candidates_from_matches
from authoroved_core.nlp.stanza_adapter import convert_document
from authoroved_core.ui.text_view import display_text_and_positions


def raw_match(offset=0, length=1):
    return {"offset": offset, "length": length, "message": "Возможная ошибка",
            "rule": {"id": "TEST", "category": {"id": "TYPOS"}}, "replacements": []}


def test_txt_preserves_bytes_whitespace_and_newlines(tmp_path):
    data = b'\xef\xbb\xbf' + 'Текст.\r\n\r\n  Пример!\n'.encode()
    path = tmp_path / "test.txt"
    path.write_bytes(data)
    doc = load_document(path)
    assert doc.original_bytes == data
    assert doc.text == 'Текст.\r\n\r\n  Пример!\n'
    assert doc.file_sha256 == sha256_bytes(data)
    assert doc.text_sha256 == sha256_bytes(doc.text.encode())


def test_encoding_must_be_explicit(tmp_path):
    path = tmp_path / "text.txt"
    path.write_bytes("Текст".encode("cp1251"))
    with pytest.raises(EncodingChoiceRequired):
        load_document(path)
    assert load_document(path, "cp1251").text == "Текст"


def test_docx_order_and_empty_paragraphs(tmp_path):
    from docx import Document
    source = Document()
    source.add_paragraph("Начало")
    source.add_paragraph("")
    table = source.add_table(rows=1, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Один", "Два"
    source.add_paragraph("Конец")
    path = tmp_path / "test.docx"
    source.save(path)
    assert load_document(path).text == "Начало\n\nОдин\tДва\nКонец"


@pytest.mark.parametrize("text", ["", "   ", "\r\n"])
def test_empty_import_rejected(tmp_path, text):
    path = tmp_path / "empty.txt"
    path.write_text(text, encoding="utf-8")
    with pytest.raises(ValueError):
        load_document(path)


def test_sha_is_deterministic_and_strict():
    assert sha256_bytes(b"abc") == "ba7816bf8f01cfea414140de5dae2223b00361a396177a9cb410ff61f20015ad"
    assert sha256_bytes(b"a\r\n") != sha256_bytes(b"a\n")


@pytest.mark.parametrize("text,start,length,fragment", [
    ("😀 ошибка", 3, 6, "ошибка"), ("😀😀 ошибка", 5, 6, "ошибка"),
    ("😀😀 ошибка", 0, 4, "😀😀"), ("а\r\nошибка", 3, 6, "ошибка"),
    ("Текст", 5, 0, ""),
])
def test_lt_utf16_spans(text, start, length, fragment):
    candidates = candidates_from_matches("d", text, [raw_match(start, length)])
    assert candidates[0].fragment == fragment
    assert text[candidates[0].span.start:candidates[0].span.end] == fragment


@pytest.mark.parametrize("start,length", [(-1, 1), (3, 1), (1, 1), (0, -1)])
def test_lt_invalid_spans_rejected(start, length):
    with pytest.raises(ValueError):
        candidates_from_matches("d", "😀", [raw_match(start, length)])


def test_unknown_languagetool_word_is_not_declared_an_error():
    match = raw_match(0, 8)
    match["rule"]["id"] = "MORFOLOGIK_RULE_RU_RU"
    match["replacements"] = [{"value": "ваххабиты"}]
    candidate = candidates_from_matches("d", "ваххабит", [match])[0]

    assert candidate.category == "Слово не распознано словарём"
    assert "ещё не означает ошибку" in candidate.explanation
    assert candidate.replacements == ("ваххабиты",)


def test_review_changes_only_decision():
    candidate = Candidate("c", "d", "Ошибка", "Орфография", "Проверить", "слво", Span(0, 4), "RULE")
    candidate.review(ReviewStatus.ACCEPTED, "Подтверждено")
    assert candidate.status == ReviewStatus.ACCEPTED
    candidate.review(ReviewStatus.REJECTED)
    assert candidate.comment == "Подтверждено"
    assert candidate.span == Span(0, 4)
    assert candidate.fragment == "слво"


def test_stanza_keeps_numeric_pos_and_never_invents_offset():
    word = NS(text="123", lemma="123", upos="NUM", feats="Case=Nom", deprel="root", head=0, id=1,
              start_char=None, end_char=None)
    token = NS(words=[word], start_char=0, end_char=3)
    doc = NS(sentences=[NS(tokens=[token])])
    result = convert_document(doc, "123")
    assert result[0].pos == "NUM" and result[0].span == Span(0, 3)
    token.end_char = 2
    assert convert_document(doc, "123")[0].span is None


def test_qt_positions_include_crlf_and_astral_symbols():
    display, positions = display_text_and_positions("😀\r\nслово")
    assert display == "😀\nслово"
    assert positions == [0, 2, 3, 3, 4, 5, 6, 7, 8]


def test_transparent_metrics():
    text = "Он дома. Он?"
    tokens = [Token("Он", "он", "PRON", {}, "nsubj", 2, 0, 1, Span(0, 2)),
              Token("дома", "дома", "ADV", {}, "root", 0, 0, 2, Span(3, 7)),
              Token(".", ".", "PUNCT", {}, "punct", 2, 0, 3, Span(7, 8)),
              Token("Он", "он", "PRON", {}, "root", 0, 1, 1, Span(9, 11)),
              Token("?", "?", "PUNCT", {}, "punct", 1, 1, 2, Span(11, 12))]
    result = {m.name: m for m in calculate_metrics(text, tokens)}
    assert result["Слова"].value == "3"
    assert result["Уникальные словоформы"].value == "2"
    assert result["Средняя длина предложения"].value == "1,5 слова"
    assert result["Средняя длина предложения"].spans == ()
    assert result["Предложения с вопросительным знаком"].value == "1"
    assert calculate_metrics(text, tokens) == calculate_metrics(text, tokens)


def test_dependency_metrics_are_explicitly_marked_as_ud_model_output():
    tokens = [Token("Он", "он", "PRON", {}, "nsubj", 2, 0, 1, Span(0, 2)),
              Token("дома", "дома", "ADV", {}, "root", 0, 0, 2, Span(3, 7))]
    metrics = {m.name: m for m in calculate_metrics("Он дома", tokens)}

    subject = metrics["подлежащее (код Stanza: nsubj)"]
    assert subject.group == "Служебная синтаксическая разметка Stanza"
    assert "не самостоятельное понятие традиционного русского синтаксиса" in subject.explanation
    assert "Universal Dependencies v2" in subject.explanation


def test_zero_word_denominators():
    assert calculate_metrics("!!!", [Token("!!!", "!!!", "PUNCT", {}, "root", 0, 0, 1, Span(0, 3))])


def test_no_legacy_imports():
    import ast
    for path in Path("authoroved_core").rglob("*.py"):
        if any(part in {".venv", "tests"} for part in path.parts):
            continue
        tree = ast.parse(path.read_text(encoding="utf-8"))
        for node in ast.walk(tree):
            if isinstance(node, ast.ImportFrom):
                assert (node.module or "").split(".")[0] not in {"analyzer", "protocol", "ui2", "expert_core"}
            if isinstance(node, ast.Import):
                assert all(alias.name.split(".")[0] not in {"analyzer", "protocol", "ui2", "expert_core"} for alias in node.names)
