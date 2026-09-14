"""Экспорт проекта исследования и автономного пакета проверки."""
from __future__ import annotations

from dataclasses import asdict
from datetime import datetime, timezone
from hashlib import sha256
import json
import os
from pathlib import Path
import re
from typing import Any
from uuid import uuid4
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document as WordDocument
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from authoroved_core import __version__
from authoroved_core.core.case_repository import CaseData, CaseRepository
from authoroved_core.core.comparison import ComparisonResult, compare_results
from authoroved_core.core.models import ReviewStatus


REPORT_TITLE = "Проект материалов исследования"
REPORT_NOTICE = (
    "Документ фиксирует автоматические измерения и решения эксперта по наблюдениям. "
    "Он не содержит вывода о тождестве или различии авторов. Значимость каждого "
    "показателя оценивает эксперт с учётом пригодности и сопоставимости материалов."
)
PACKAGE_FORMAT = "authoroved-verification-package"
PACKAGE_VERSION = 1
BLUE = "173B70"
PALE_BLUE = "EAF1F8"
LIGHT_GRAY = "D9D9D9"


class ReportError(ValueError):
    """Материалов недостаточно для формирования проекта."""


def _json_bytes(value: Any) -> bytes:
    return json.dumps(value, ensure_ascii=False, sort_keys=True, indent=2).encode("utf-8")


def _atomic_replace(path: Path, writer) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(f".{path.name}.{uuid4().hex}.tmp")
    try:
        writer(temporary)
        if not temporary.is_file() or temporary.stat().st_size == 0:
            raise OSError("Экспорт не создал непустой файл.")
        os.replace(temporary, path)
    finally:
        if temporary.exists():
            temporary.unlink()
    return path


def _require_ready(case: CaseData) -> None:
    if len(case.materials) != 2 or len(case.results) != 2 or not all(case.materials) or not all(case.results):
        raise ReportError("Для экспорта нужно загрузить и проанализировать два текста.")


def _metric_value(result, name: str, default="Нет данных") -> str:
    metric = next((item for item in result.metrics if item.name == name), None)
    return metric.value if metric else default


def _safe_name(name: str) -> str:
    clean = re.sub(r"[^\w.() -]+", "_", Path(name).name, flags=re.UNICODE).strip(" .")
    return clean or "document.bin"


def _component_rows(case: CaseData) -> list[dict[str, str]]:
    rows = [{"component": "Авторовед Core", "version": __version__, "details": "локальная обработка"}]
    seen = set()
    for result in case.results:
        for key, title in (("stanza", "Stanza"), ("languagetool", "LanguageTool")):
            metadata = result.metadata.get(key) if result else None
            if not isinstance(metadata, dict):
                continue
            version = str(metadata.get("version") or "не указана")
            details = []
            if key == "stanza":
                details.extend([
                    f"устройство: {metadata.get('device', 'не указано')}",
                    f"PyTorch: {metadata.get('torch_version', 'не указана')}",
                ])
            else:
                details.extend([
                    "режим: локальный командный запуск" if metadata.get("mode") == "local-cli"
                    else f"режим: {metadata.get('mode', 'не указан')}",
                    f"сборка: {metadata.get('build_date') or 'не указана'}",
                ])
            identity = (title, version, tuple(details))
            if identity not in seen:
                seen.add(identity)
                rows.append({"component": title, "version": version, "details": "; ".join(details)})
    return rows


def _material_rows(case: CaseData) -> list[dict[str, Any]]:
    rows = []
    for index, (material, result) in enumerate(zip(case.materials, case.results), 1):
        rows.append({
            "number": index,
            "name": material.name,
            "size_bytes": len(material.original_bytes),
            "file_sha256": material.file_sha256,
            "text_sha256": material.text_sha256,
            "encoding": material.encoding,
            "words": _metric_value(result, "Слова"),
            "sentences": _metric_value(result, "Предложения"),
            "paragraphs": _metric_value(result, "Абзацы"),
            "analysis_errors": list(result.errors),
        })
    return rows


def _review_rows(case: CaseData) -> list[dict[str, Any]]:
    rows = []
    for index, result in enumerate(case.results, 1):
        counts = {status: sum(item.status == status for item in result.candidates) for status in ReviewStatus}
        rows.append({
            "text": index,
            "total": len(result.candidates),
            "accepted": counts[ReviewStatus.ACCEPTED],
            "rejected": counts[ReviewStatus.REJECTED],
            "unreviewed": counts[ReviewStatus.NEW],
        })
    return rows


def _comparison_data(comparison: ComparisonResult) -> dict[str, Any]:
    return {
        "limitations": list(comparison.limitations),
        "metrics": [
            {
                "name": item.name, "group": item.group,
                "text_1": item.value_first, "text_2": item.value_second,
                "difference": item.difference, "direction": item.direction,
                "basis": item.basis, "explanation": item.explanation,
                "limitation": item.caution,
            } for item in comparison.metrics
        ],
        "accepted_observations": [
            {
                "category": item.category, "label": item.label,
                "text_1_count": item.count_first, "text_2_count": item.count_second,
                "relation": item.relation, "grouping_basis": item.basis,
            } for item in comparison.accepted_groups
        ],
    }


def _feature_data(case: CaseData) -> list[dict[str, Any]]:
    rows = []
    for slot, result in enumerate(case.results, 1):
        if result is None:
            continue
        for item in result.feature_observations:
            rows.append({
                "text": slot,
                "feature_id": item.feature_id,
                "raw_value": item.raw_value,
                "normalized_value": item.normalized_value,
                "applicability": item.applicability.value,
                "method_version": item.method_version,
                "source": [asdict(source) for source in item.source],
                "confidence": item.confidence,
                "expert_status": item.expert_status.value,
                "expert_value": item.expert_value,
                "expert_comment": item.expert_comment,
                "limitations": list(item.limitations),
                "evidence": [
                    {"quote": evidence.quote, "start": evidence.span.start,
                     "end": evidence.span.end, "label": evidence.label}
                    for evidence in item.evidence
                ],
            })
    return rows


def _set_font(run, name="Arial", size=11, bold=None, color="000000"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def _shade(cell, fill: str):
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def _set_cell_margins(cell, top=90, start=100, bottom=90, end=100):
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _set_table_borders(table):
    properties = table._tbl.tblPr
    borders = properties.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:color"), LIGHT_GRAY)


def _repeat_header(row):
    properties = row._tr.get_or_add_trPr()
    marker = OxmlElement("w:tblHeader")
    marker.set(qn("w:val"), "true")
    properties.append(marker)


def _dont_split(row):
    properties = row._tr.get_or_add_trPr()
    marker = OxmlElement("w:cantSplit")
    marker.set(qn("w:val"), "true")
    properties.append(marker)


def _table(document, headers: list[str], rows: list[list[str]], widths: list[float] | None = None):
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    _set_table_borders(table)
    _repeat_header(table.rows[0])
    _dont_split(table.rows[0])
    for index, text in enumerate(headers):
        cell = table.rows[0].cells[index]
        _shade(cell, BLUE)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        _set_cell_margins(cell)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _set_font(paragraph.add_run(text), size=9, bold=True, color="FFFFFF")
        if widths:
            cell.width = Inches(widths[index])
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        _dont_split(table.rows[-1])
        for index, value in enumerate(values):
            cell = cells[index]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            _set_cell_margins(cell)
            if row_index % 2:
                _shade(cell, PALE_BLUE)
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if index and len(str(value)) < 35 else WD_ALIGN_PARAGRAPH.LEFT
            _set_font(paragraph.add_run(str(value)), size=9)
            if widths:
                cell.width = Inches(widths[index])
    document.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def _configure_document(document):
    section = document.sections[0]
    _configure_section(section)
    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08
    for style_name, size in (("Title", 22), ("Heading 1", 15), ("Heading 2", 12)):
        style = document.styles[style_name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(6)
        paragraph_properties = style._element.get_or_add_pPr()
        borders = paragraph_properties.find(qn("w:pBdr"))
        if borders is not None:
            paragraph_properties.remove(borders)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_font(footer.add_run("Авторовед Core · проект материалов исследования"), size=8, color="666666")


def _configure_section(section):
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)


def _add_heading(document, text, level=1):
    paragraph = document.add_heading(text, level=level)
    paragraph.paragraph_format.keep_with_next = True
    return paragraph


def _new_report_page(document):
    """Начинает страницу секционным разрывом, устойчивым при экспорте Word в PDF."""
    section = document.add_section(WD_SECTION.NEW_PAGE)
    _configure_section(section)
    section.footer.is_linked_to_previous = True


def _add_report_body(document, case: CaseData, comparison: ComparisonResult):
    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title_properties = title._p.get_or_add_pPr()
    borders = title_properties.find(qn("w:pBdr"))
    if borders is not None:
        title_properties.remove(borders)
    _set_font(title.add_run(REPORT_TITLE), size=22, bold=True)
    subtitle = document.add_paragraph()
    _set_font(subtitle.add_run(f"Дело {case.case_id}"), size=10, color="555555")
    document.add_paragraph(REPORT_NOTICE)

    _add_heading(document, "Объекты исследования")
    material_rows = _material_rows(case)
    _table(document, ["Объект", "Файл", "Объём", "Контрольная сумма SHA-256"], [
        [f"Текст {row['number']}", row["name"],
         f"{row['size_bytes']} байт\n{row['words']} слов\n{row['sentences']} предложений",
         row["file_sha256"]]
        for row in material_rows
    ], [0.65, 1.55, 1.25, 3.55])
    document.add_paragraph(
        "Контрольная сумма относится к исходным байтам файла. Извлечённый текст "
        "контролируется отдельной суммой и хранится в зашифрованном деле."
    )

    _add_heading(document, "Условия автоматизированной обработки")
    _table(document, ["Компонент", "Версия", "Условия"], [
        [row["component"], row["version"], row["details"]] for row in _component_rows(case)
    ], [1.55, 1.1, 4.35])
    partial = any(result.errors for result in case.results)
    document.add_paragraph(
        "Обработка выполнена локально. Автоматическая разметка может содержать ошибки. "
        + ("В одном или обоих текстах часть анализаторов завершилась с ошибкой."
           if partial else "Оба результата не содержат сообщений о сбое анализаторов.")
    )

    _add_heading(document, "Проверка автоматических рекомендаций")
    _table(document, ["Материал", "Всего", "Принято", "Не учтено", "Не рассмотрено"], [
        [f"Текст {row['text']}", row["total"], row["accepted"], row["rejected"], row["unreviewed"]]
        for row in _review_rows(case)
    ], [1.4, 1.0, 1.1, 1.35, 1.55])
    document.add_paragraph(
        "В сопоставление включены только наблюдения, которые эксперт сохранил. "
        "Отклонённые и нерассмотренные рекомендации в доказательную часть не переносятся."
    )

    _new_report_page(document)
    _add_heading(document, "Сопоставление измерений")
    grouped: dict[str, list] = {}
    for item in comparison.metrics:
        grouped.setdefault(item.group, []).append(item)
    first_comparison_block = True
    for group, items in grouped.items():
        # Word sometimes loses the text of a repeated table header after an
        # automatic page break.  Small, explicit blocks keep every header
        # visible and also keep a data row from being divided between pages.
        # Пять строк гарантированно помещаются на страницу Letter вместе с
        # заголовком даже при длинных пояснениях в последнем столбце. Более
        # крупный блок Word иногда пытался удержать целиком и вытеснял его
        # верхнюю часть за печатное поле.
        block_count = max(1, (len(items) + 4) // 5)
        block_size = (len(items) + block_count - 1) // block_count
        for offset in range(0, len(items), block_size):
            block = items[offset:offset + block_size]
            heading_text = group if offset == 0 else f"{group} — продолжение"
            if not first_comparison_block:
                _new_report_page(document)
            _add_heading(document, heading_text, level=2)
            first_comparison_block = False
            _table(document, ["Показатель", "Текст 1", "Текст 2", "Разница", "Основа сопоставления"], [
                [item.name, item.value_first, item.value_second,
                 f"{item.difference}; {item.direction}", f"{item.basis}. {item.caution}"]
                for item in block
            ], [1.6, 1.0, 1.0, 1.05, 2.35])

    _new_report_page(document)
    _add_heading(document, "Принятые наблюдения")
    if comparison.accepted_groups:
        accepted = list(comparison.accepted_groups)
        for offset in range(0, len(accepted), 8):
            if offset:
                _new_report_page(document)
                _add_heading(document, "Принятые наблюдения — продолжение")
            block = accepted[offset:offset + 8]
            _table(document, ["Наблюдение", "Текст 1", "Текст 2", "Результат", "Основа группировки"], [
                [item.label, item.count_first, item.count_second, item.relation, item.basis]
                for item in block
            ], [1.9, 0.7, 0.7, 1.25, 2.45])
    else:
        document.add_paragraph("Эксперт не сохранил наблюдений для сопоставления.")

    _add_heading(document, "Ограничения")
    for item in comparison.limitations:
        document.add_paragraph(item, style="List Bullet")
    document.add_paragraph(
        "Разницы не имеют автоматических порогов доказательственной значимости. "
        "Формулирование экспертного вывода в этот проект не входит."
    )

    _add_heading(document, "Журнал и целостность")
    document.add_paragraph(
        f"В журнале дела {len(case.audit)} записей. Последняя запись имеет контрольную сумму "
        f"{case.audit[-1].entry_hash if case.audit else 'отсутствует'}. "
        "Целостность материалов и цепочки журнала проверена перед экспортом."
    )


class ReportService:
    """Формирует воспроизводимые выгрузки без автоматического вывода об авторстве."""

    def __init__(self, case_repository: CaseRepository | None = None):
        self.case_repository = case_repository or CaseRepository()

    def _prepare(self, case: CaseData, comparison: ComparisonResult | None) -> ComparisonResult:
        _require_ready(case)
        self.case_repository.verify_integrity(case)
        return comparison or compare_results(case.results[0], case.results[1])

    def export_docx(self, path: str | Path, case: CaseData,
                    comparison: ComparisonResult | None = None) -> Path:
        path = Path(path)
        if path.suffix.lower() != ".docx":
            path = path.with_suffix(".docx")
        comparison = self._prepare(case, comparison)
        document = WordDocument()
        _configure_document(document)
        _add_report_body(document, case, comparison)
        return _atomic_replace(path, lambda temporary: document.save(temporary))

    def export_verification_package(self, path: str | Path, case: CaseData,
                                    comparison: ComparisonResult | None = None,
                                    *, include_sources: bool = False) -> Path:
        path = Path(path)
        if path.suffix.lower() != ".zip":
            path = path.with_suffix(".zip")
        comparison = self._prepare(case, comparison)
        generated_at = datetime.now(timezone.utc).isoformat()
        entries: dict[str, bytes] = {
            "audit.json": _json_bytes([asdict(item) for item in case.audit]),
            "components.json": _json_bytes(_component_rows(case)),
            "materials.json": _json_bytes(_material_rows(case)),
            "comparison.json": _json_bytes(_comparison_data(comparison)),
            "features.json": _json_bytes(_feature_data(case)),
        }
        if include_sources:
            for index, material in enumerate(case.materials, 1):
                entries[f"sources/{index}_{_safe_name(material.name)}"] = material.original_bytes
                entries[f"sources/{index}_extracted.txt"] = material.text.encode("utf-8")
        manifest = {
            "format": PACKAGE_FORMAT,
            "version": PACKAGE_VERSION,
            "case_id": case.case_id,
            "generated_at": generated_at,
            "program_version": __version__,
            "source_documents_included": include_sources,
            "contents": sorted([*entries, "manifest.json", "checksums.json"]),
            "notice": REPORT_NOTICE,
        }
        entries["manifest.json"] = _json_bytes(manifest)
        checksums = {name: sha256(data).hexdigest() for name, data in sorted(entries.items())}
        entries["checksums.json"] = _json_bytes({"algorithm": "SHA-256", "files": checksums})

        def write(temporary: Path):
            with ZipFile(temporary, "w", ZIP_DEFLATED) as archive:
                for name, data in sorted(entries.items()):
                    archive.writestr(name, data)

        return _atomic_replace(path, write)

    @staticmethod
    def verify_verification_package(path: str | Path) -> None:
        with ZipFile(path, "r") as archive:
            names = set(archive.namelist())
            if "checksums.json" not in names or "manifest.json" not in names:
                raise ReportError("В проверочном пакете отсутствуют служебные файлы.")
            manifest = json.loads(archive.read("manifest.json").decode("utf-8"))
            checksums = json.loads(archive.read("checksums.json").decode("utf-8"))
            if manifest.get("format") != PACKAGE_FORMAT or manifest.get("version") != PACKAGE_VERSION:
                raise ReportError("Формат проверочного пакета не поддерживается.")
            if set(manifest.get("contents", [])) != names:
                raise ReportError("Состав проверочного пакета не совпадает с манифестом.")
            for name, expected in checksums.get("files", {}).items():
                if name not in names or sha256(archive.read(name)).hexdigest() != expected:
                    raise ReportError(f"Не совпадает контрольная сумма файла {name}.")
