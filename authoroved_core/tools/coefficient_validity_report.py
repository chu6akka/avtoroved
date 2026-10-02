"""Таблицы и графики для диплома по отчёту `coefficient_validity.json`.

Создаёт DOCX с описанием материала и метода, таблицами по 40 классическим
коэффициентам, лучшим показателям и минимальному объёму текста, и два графика PNG.
Выводов об авторстве документ не содержит.
"""
from __future__ import annotations

import argparse
import json
import os
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor

CLASSIC_PREFIXES = ("Индекс ", "Коэффициент ")
CHART_COLORS = ("#4356c8", "#d9822b", "#2f8f5b", "#b5446e", "#6b5fb5", "#8a8a8a")
_QT_APP = None


def number(value, digits=2) -> str:
    if value is None:
        return "—"
    return f"{value:.{digits}f}".replace(".", ",")


def p_text(value) -> str:
    if value is None:
        return "—"
    return "< 0,001" if value < 0.001 else number(value, 3)


def short(name: str, limit: int = 70) -> str:
    return name if len(name) <= limit else name[:limit - 1] + "…"


def set_cell(cell, text, bold=False, size=9, align=None):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text)
    run.font.size = Pt(size)
    run.bold = bold
    if align:
        paragraph.alignment = align


def table(document, header, rows, widths=None):
    grid = document.add_table(rows=1, cols=len(header))
    grid.style = "Table Grid"
    for cell, text in zip(grid.rows[0].cells, header):
        set_cell(cell, text, bold=True)
    for row in rows:
        cells = grid.add_row().cells
        for index, (cell, text) in enumerate(zip(cells, row)):
            set_cell(cell, str(text), align=None if index == 0 else WD_ALIGN_PARAGRAPH.CENTER)
    if widths:
        for row in grid.rows:
            for cell, width in zip(row.cells, widths):
                cell.width = Cm(width)
    return grid


def classic_order(name: str) -> tuple:
    """Индексы 1–20, затем коэффициенты 1–20, затем дополнительные именованные индексы."""
    head = name.split(".")[0]
    digits = head.split(" ")[-1] if head.split(" ")[-1].isdigit() else ""
    if not digits:
        return 2, name
    return (0 if name.startswith("Индекс ") else 1), f"{int(digits):03d}"


def draw_chart(path: Path, title: str, series: dict[str, list[tuple[int, float]]],
               y_label: str, y_range: tuple[float, float], reference: float | None = None,
               y_digits: int = 2):
    """Линейный график на QPainter (программа уже содержит Qt; matplotlib не нужен)."""
    os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
    from PyQt6.QtCore import QPointF, QRectF, Qt
    from PyQt6.QtGui import QColor, QFont, QImage, QPainter, QPen
    from PyQt6.QtWidgets import QApplication
    from authoroved_core.ui.branding import load_fonts

    global _QT_APP
    # Ссылку нужно держать: без неё объект приложения сразу удаляется и рисование падает.
    _QT_APP = QApplication.instance() or QApplication([])
    load_fonts()
    width, height = 1600, 900
    image = QImage(width, height, QImage.Format.Format_ARGB32)
    image.fill(QColor("#ffffff"))
    painter = QPainter(image)
    painter.setRenderHint(QPainter.RenderHint.Antialiasing)
    left, right, top, bottom = 140, 470, 90, 120
    plot = QRectF(left, top, width - left - right, height - top - bottom)
    xs = sorted({x for points in series.values() for x, _ in points})
    y_low, y_high = y_range

    def to_point(x, y):
        fx = (xs.index(x) / (len(xs) - 1)) if len(xs) > 1 else 0.5
        fy = (y - y_low) / (y_high - y_low)
        return QPointF(plot.left() + fx * plot.width(), plot.bottom() - fy * plot.height())

    painter.setFont(QFont("Golos Text", 22, QFont.Weight.DemiBold))
    painter.setPen(QColor("#161d33"))
    painter.drawText(QRectF(left, 20, width - left, 50), Qt.AlignmentFlag.AlignLeft, title)
    painter.setFont(QFont("Golos Text", 15))
    grid_pen = QPen(QColor("#e3e6ee"), 1)
    for step in range(6):
        y = y_low + (y_high - y_low) * step / 5
        point = to_point(xs[0], y)
        painter.setPen(grid_pen)
        painter.drawLine(QPointF(plot.left(), point.y()), QPointF(plot.right(), point.y()))
        painter.setPen(QColor("#556079"))
        painter.drawText(QRectF(10, point.y() - 14, left - 20, 28),
                         Qt.AlignmentFlag.AlignRight | Qt.AlignmentFlag.AlignVCenter, number(y, y_digits))
    for x in xs:
        point = to_point(x, y_low)
        painter.drawText(QRectF(point.x() - 50, plot.bottom() + 10, 100, 30), Qt.AlignmentFlag.AlignHCenter, str(x))
    painter.drawText(QRectF(plot.left(), height - 60, plot.width(), 40), Qt.AlignmentFlag.AlignHCenter,
                     "объём фрагмента, слов")
    painter.save()
    painter.translate(30, plot.center().y())
    painter.rotate(-90)
    painter.drawText(QRectF(-200, -20, 400, 40), Qt.AlignmentFlag.AlignHCenter, y_label)
    painter.restore()
    if reference is not None:
        point = to_point(xs[0], reference)
        painter.setPen(QPen(QColor("#9aa3bd"), 2, Qt.PenStyle.DashLine))
        painter.drawLine(QPointF(plot.left(), point.y()), QPointF(plot.right(), point.y()))
    for index, (name, points) in enumerate(series.items()):
        color = QColor(CHART_COLORS[index % len(CHART_COLORS)])
        painter.setPen(QPen(color, 4))
        chain = [to_point(x, y) for x, y in sorted(points) if y is not None]
        for a, b in zip(chain, chain[1:]):
            painter.drawLine(a, b)
        painter.setBrush(color)
        for point in chain:
            painter.drawEllipse(point, 6, 6)
        legend_y = top + 10 + index * 70
        painter.drawLine(QPointF(width - right + 30, legend_y + 14), QPointF(width - right + 70, legend_y + 14))
        painter.setPen(QColor("#161d33"))
        painter.drawText(QRectF(width - right + 82, legend_y - 4, right - 95, 64),
                         Qt.TextFlag.TextWordWrap, short(name, 60))
    painter.end()
    image.save(str(path))


def build(report: dict, output: Path, charts_dir: Path, combination: dict | None = None) -> Path:
    first, second = report["discrimination"], report["length"]
    metrics = first["by_metric"]
    length_metrics = second["by_metric"]
    document = Document()
    style = document.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)

    document.add_heading("Проверка пригодности морфологических коэффициентов", level=1)
    note = document.add_paragraph(
        "Проверяется пригодность признаков, а не точность установления автора программой. "
        "Принадлежность текстов одному автору в корпусе определена по учётной записи Pikabu "
        "и не равнозначна процессуально удостоверенному авторству. Порогов и выводов о "
        "тождестве автора из этих данных не следует.")
    note.runs[0].italic = True

    document.add_heading("1. Материал и метод", level=2)
    document.add_paragraph(
        f"Корпус Pilot 02c: {first['authors']} авторов, по 4 основных текста "
        f"({first['texts']} текстов) и по 2 резервных, публикации Pikabu объёмом 300–700 слов "
        "(медиана около 466). Морфологическая разметка — Stanza 1.14 (модели SynTagRus), "
        "показатели — те же, что выводит программа «Авторовед».")
    document.add_paragraph(
        "Различительная сила показателя оценивалась мерой AUC: вероятностью того, что у двух "
        "текстов разных авторов значения показателя расходятся сильнее, чем у двух текстов "
        "одного автора. AUC = 0,5 означает, что показатель не несёт сведений об авторе; 1,0 — "
        "полное разделение. Значимость проверялась перестановочным тестом "
        f"({first['permutations']} случайных перестановок меток автора между текстами): он "
        "учитывает, что один текст входит во многие пары. Поправка на множественность "
        "сравнений — по Бенджамини–Хохбергу (q). ICC — доля разброса значений, объясняемая "
        "автором. Чтобы исключить влияние темы, AUC пересчитывался только на парах текстов "
        f"без единой общей метки Pikabu (таких пар {number(first['topic_neutral_pairs_share'] * 100, 0)} %).")
    document.add_paragraph(
        "Влияние объёма проверялось на фрагментах: тексты каждого автора объединялись "
        f"(все 6, около 2800 слов), из них {second['repetitions']} раз случайно извлекались по два "
        f"непересекающихся фрагмента длиной {', '.join(map(str, second['lengths']))} слов. Для "
        "каждого объёма считались AUC фрагментов и отклонение — медиана относительной разницы "
        "между значением на фрагменте и на всём материале автора. Минимальный объём — "
        "наименьшая длина, начиная с которой отклонение не превышает 20 % (10 %).")
    document.add_paragraph(
        "Оценка «различает авторов»: q < 0,05, AUC ≥ 0,6 и значимость сохраняется на парах без "
        "общих тем; «слабо различает»: значимо, но AUC < 0,6; «различает, но не без общих тем»: "
        "значимость исчезает при исключении общих тем; «не отличается от случайного»: q ≥ 0,05.")

    document.add_heading("2. Классические коэффициенты", level=2)
    document.add_paragraph("Таблица 1. Индексы идиостиля (Соколова Т. П.), морфологические коэффициенты "
                           "(Вул С. М., Галяшина Е. И.) и дополнительные стилеметрические индексы")
    classic = sorted((name for name in metrics if name.startswith(CLASSIC_PREFIXES)), key=classic_order)
    rows = []
    for name in classic:
        item = metrics[name]
        rows.append([short(name, 62), number(item["auc"]), number(item["icc"]), p_text(item["q"]),
                     number(item["auc_topic_neutral"]), item["verdict"]])
    table(document, ["Коэффициент", "AUC", "ICC", "q", "AUC без общих тем", "Оценка"], rows,
          widths=[7.2, 1.3, 1.3, 1.6, 1.8, 3.6])
    excluded = [number_ for number_ in range(1, 21) if not any(
        name.startswith(f"Индекс {number_:02d}.") for name in metrics)]
    excluded_sae = [number_ for number_ in range(1, 21) if not any(
        name.startswith(f"Коэффициент {number_:02d}.") for name in metrics)]
    notes = []
    if 7 in excluded:
        notes.append("индекс 7 (абстрактные / конкретные существительные) требует семантической "
                     "разметки и автоматически не рассчитывается")
    for number_ in [n for n in excluded if n != 7]:
        notes.append(f"индекс {number_} не определён более чем у 20 % текстов")
    for number_ in excluded_sae:
        notes.append(f"коэффициент {number_} не определён более чем у 20 % текстов (нулевой знаменатель; "
                     "для коэффициента 19 — тексты без деепричастий)")
    if notes:
        document.add_paragraph("В проверку не вошли: " + "; ".join(notes) + ".")

    document.add_heading("3. Наиболее различающие показатели", level=2)
    document.add_paragraph("Таблица 2. Пятнадцать показателей с наибольшим AUC среди всех рассчитываемых "
                           f"программой ({len(metrics)} показателей)")
    top = list(metrics.items())[:15]
    table(document, ["Показатель", "Раздел", "AUC", "ICC", "AUC без общих тем", "Оценка"],
          [[short(name, 55), short(item["group"], 28), number(item["auc"]), number(item["icc"]),
            number(item["auc_topic_neutral"]), item["verdict"]] for name, item in top],
          widths=[6.0, 3.4, 1.3, 1.3, 1.8, 3.2])
    counts = {}
    for item in metrics.values():
        counts[item["verdict"]] = counts.get(item["verdict"], 0) + 1
    document.add_paragraph("Распределение оценок по всем показателям: " + "; ".join(
        f"{key} — {value}" for key, value in sorted(counts.items(), key=lambda kv: -kv[1])) + ".")

    document.add_heading("4. Минимальный объём текста", level=2)
    document.add_paragraph("Таблица 3. Отклонение значения на фрагменте от значения на всём материале "
                           "автора и минимальный объём для классических коэффициентов")
    lengths = [str(value) for value in second["lengths"]]
    shown = [value for value in ("200", "300", "500", "1000") if value in lengths]
    rows = []
    for name in classic:
        item = length_metrics.get(name)
        if not item:
            continue
        by_length = item["by_length"]

        def deviation(value):
            found = (by_length.get(value) or {}).get("deviation")
            return "—" if found is None else number(found * 100, 0) + " %"

        rows.append([short(name, 55)] + [deviation(value) for value in shown]
                    + [str(item["min_words_20"] or "> " + lengths[-1]),
                       str(item["min_words_10"] or "> " + lengths[-1])])
    table(document, ["Коэффициент"] + [f"откл. {value} сл." for value in shown] + ["мин. объём (20 %)", "мин. объём (10 %)"],
          rows, widths=[6.4] + [1.6] * len(shown) + [2.0, 2.0])

    charts_dir.mkdir(parents=True, exist_ok=True)
    best = [name for name in classic if metrics[name]["verdict"] == "различает авторов"]
    best = sorted(best, key=lambda name: -(metrics[name]["auc"] or 0))[:5] or classic[:5]
    auc_series = {name: [(int(k), v["auc"]) for k, v in length_metrics[name]["by_length"].items()]
                  for name in best if name in length_metrics}
    deviation_series = {name: [(int(k), (v["deviation"] or 0) * 100) for k, v in length_metrics[name]["by_length"].items()]
                        for name in best if name in length_metrics}
    auc_chart = charts_dir / "auc_by_length.png"
    deviation_chart = charts_dir / "deviation_by_length.png"
    draw_chart(auc_chart, "Различительная сила (AUC) в зависимости от объёма фрагмента", auc_series,
               "AUC", (0.5, 0.75), reference=None)
    max_deviation = max((y for points in deviation_series.values() for _, y in points), default=50)
    draw_chart(deviation_chart, "Отклонение от значения на всём материале автора, %", deviation_series,
               "отклонение, %", (0, max(25.0, round(max_deviation / 10 + 1) * 10)), reference=20,
               y_digits=0)
    document.add_paragraph("Рисунок 1. AUC лучших классических коэффициентов на фрагментах разного объёма "
                           "(0,5 — отсутствие различения)")
    document.add_picture(str(auc_chart), width=Cm(16))
    document.add_paragraph("Рисунок 2. Отклонение значения на фрагменте от значения на всём материале "
                           "автора (пунктир — 20 %)")
    document.add_picture(str(deviation_chart), width=Cm(16))

    if combination:
        add_combination_section(document, combination, charts_dir)
    document.add_heading("6. Выводы" if combination else "5. Выводы", level=2)
    for text in conclusions(report, classic, combination):
        document.add_paragraph(text, style="List Number")

    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(output))
    return output


def add_combination_section(document, combination: dict, charts_dir: Path) -> None:
    texts, lengths = combination["texts"], combination["lengths"]
    methods = texts["methods"]
    document.add_heading("5. Совокупность признаков", level=2)
    document.add_paragraph(
        "Проверялось, различает ли авторов набор показателей лучше, чем самый сильный из них. "
        "Расстояние между текстами по набору — среднее модулей разностей нормированных значений; "
        "взвешенный вариант — логистическая регрессия по тем же разностям. Чтобы не подогнать "
        f"набор под данные, авторы делились на {combination['folds']} групп: отбор показателей, "
        "нормировка и веса считались на четырёх группах, AUC — на пятой, чьих текстов модель не "
        f"видела; разбиение повторялось {combination['repeats']} раз. Лучший одиночный показатель "
        "тоже выбирался только по обучающим авторам.")
    document.add_paragraph("Таблица 4. AUC наборов показателей на текстах, не участвовавших в отборе "
                           "(в скобках — 95 % разброс по разбиениям) и на фрагментах разного объёма")
    shown = sorted(lengths, key=int)
    rows = []
    for key, item in methods.items():
        rows.append([item["title"], f"{number(item['auc'])} ({number(item['low'])}–{number(item['high'])})"]
                    + [number(lengths[length]["methods"].get(key)) for length in shown])
    table(document, ["Набор", "Целые тексты"] + [f"{length} сл." for length in shown], rows,
          widths=[6.0, 2.8] + [1.5] * len(shown))
    series = {methods[key]["title"].split(" (")[0].split(",")[0]: [(int(length), lengths[length]["methods"][key])
                                                                    for length in shown]
              for key in ("single", "classic", "all", "logistic") if key in methods}
    chart = charts_dir / "combination_by_length.png"
    draw_chart(chart, "Одиночный показатель и совокупность: AUC по объёму фрагмента", series,
               "AUC", (0.45, 0.9), reference=0.5)
    document.add_paragraph("Рисунок 3. AUC лучшего одиночного показателя и наборов показателей на "
                           "фрагментах разного объёма (пунктир — 0,5)")
    document.add_picture(str(chart), width=Cm(16))
    document.add_paragraph(
        "AUC здесь — не доля верных заключений, а вероятность того, что из двух пар текстов "
        "(одного автора и разных) пара одного автора окажется ближе. Порога, по которому пару "
        "можно было бы отнести к «одному автору», из этих данных не следует.")


def conclusions(report: dict, classic: list[str], combination: dict | None = None) -> list[str]:
    """Выводы из чисел отчёта: формулировки фиксированы, числа подставляются."""
    metrics = report["discrimination"]["by_metric"]
    length = report["length"]["by_metric"]
    total = len(metrics)
    strong = [name for name, item in metrics.items() if item["verdict"] == "различает авторов"]
    best_name, best = next(iter(metrics.items()))
    topic_only = [name for name in classic if metrics[name]["verdict"] == "различает, но не без общих тем"]
    random_like = [name for name in classic if metrics[name]["verdict"] == "не отличается от случайного"]
    strong_classic = [name for name in classic if name in strong]
    lengths = report["length"]["lengths"]

    def label(name):
        """«Индекс 06. …» → «индекс 6», «Индекс местоименности: …» → «индекс местоименности»."""
        if not name.startswith(("Индекс ", "Коэффициент ")):
            return f"«{name}»"
        head = name.split(".")[0] if "." in name.split(":")[0] else name.split(":")[0]
        kind, _, rest = head.partition(" ")
        return f"{kind.lower()} {rest.lstrip('0') if rest.isdigit() else rest}"

    def listed(names):
        return ", ".join(label(name) for name in names)

    best_length = length.get(best_name, {}).get("by_length", {})
    smallest, largest = str(lengths[0]), str(lengths[-1])
    items = [
        f"Ни один показатель в отдельности не разделяет авторов надёжно: наибольший AUC — "
        f"{number(best['auc'])} («{best_name}»). Из {total} показателей строгому критерию отвечают "
        f"{len(strong)}, из них классических коэффициентов — {len(strong_classic)}: "
        f"{listed(strong_classic)}. Морфологические коэффициенты "
        "пригодны как совокупность общих признаков, а не как самостоятельное основание вывода.",
    ]
    if topic_only:
        items.append(
            "На парах текстов без общих тем различие исчезает у следующих показателей: "
            f"{listed(topic_only)}. Они отражают тему и жанр текста, а не автора, и при "
            "различающихся по теме материалах интерпретироваться не должны.")
    if random_like:
        items.append(f"От случайности не отличаются и сведений об авторе не несут: {listed(random_like)}.")
    if best_length.get(smallest) and best_length.get(largest):
        items.append(
            f"Различительная сила растёт с объёмом: у лучшего показателя AUC равен "
            f"{number(best_length[smallest]['auc'])} на фрагменте {smallest} слов и "
            f"{number(best_length[largest]['auc'])} на {largest} словах.")
    stable = sorted((name for name in classic if length.get(name, {}).get("min_words_10")),
                    key=lambda name: length[name]["min_words_10"])
    unstable = [name for name in classic if name in length and length[name]["min_words_20"] is None]
    if stable:
        items.append(
            f"Устойчивее всего значение показателя «{label(stable[0])}»: отклонение не превышает 10 % уже с "
            f"{length[stable[0]]['min_words_10']} слов. Для большинства классических коэффициентов "
            "допуск 20 % достигается при 200–300 словах, допуск 10 % — лишь около 1000 слов.")
    if unstable:
        items.append(
            f"Не достигают даже допуска 20 % на {largest} словах: {listed(unstable)}. Эти показатели "
            "опираются на редкие формы (деепричастия, краткие прилагательные, числительные) или, как "
            "соотношение времён глагола, зависят от способа повествования, и на текстах обычного "
            "объёма их значения неустойчивы.")
    if combination:
        methods, lengths = combination["texts"]["methods"], combination["lengths"]
        longest = max(lengths, key=int)
        items.append(
            f"Совокупность показателей различает авторов заметно лучше любого из них: на текстах, не "
            f"участвовавших в отборе, AUC лучшего одиночного показателя — {number(methods['single']['auc'])}, "
            f"среднего по всем показателям — {number(methods['all']['auc'])}, взвешенной модели — "
            f"{number(methods['logistic']['auc'])}; на фрагментах {longest} слов — "
            f"{number(lengths[longest]['methods']['single'])} против {number(lengths[longest]['methods']['all'])}. "
            "Подбор весов устойчивого выигрыша перед простым средним не даёт, а на малом материале "
            "уступает ему: для эксперта достаточно правила рассматривать показатели в совокупности.")
        leading = combination["texts"].get("logistic_leading", [])
        if any("род: женский" in name for name in leading):
            items.append(
                "Среди наиболее весомых признаков взвешенной модели — доля глаголов женского рода: в "
                "русском языке род прошедшего времени при повествовании от первого лица выдаёт пол "
                "рассказчика. Морфологические показатели, таким образом, несут и диагностические "
                "сведения об авторе, а не только сведения об идиостиле.")
    items.append(
        "Для экспертной практики: на материалах объёмом в несколько сотен слов морфологические "
        "коэффициенты допустимо приводить лишь как ориентирующие; при различии материалов по теме "
        "и жанру следует исключать тематически зависимые показатели; основание вывода должны "
        "составлять частные признаки (ошибки, навыки письма), а не отдельные коэффициенты.")
    items.append(
        "Ограничения: корпус однороден по источнику (Pikabu), принадлежность текстов автору "
        "установлена по учётной записи; авторы склонны писать в устойчивом жанре, поэтому часть "
        "«авторского» разброса может объясняться жанром. Отбор пар без общих меток снимает "
        "тематическое влияние лишь частично. Разметка Stanza автоматическая и содержит ошибки.")
    return items


def main():
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--input", type=Path, default=Path("authoroved_core/artifacts/coefficient_validity.json"))
    parser.add_argument("--output", type=Path,
                        default=Path("authoroved_core/artifacts/diploma/Проверка_коэффициентов.docx"))
    args = parser.parse_args()
    report = json.loads(args.input.read_text(encoding="utf-8"))
    combination_path = args.input.with_name("feature_combination.json")
    combination = (json.loads(combination_path.read_text(encoding="utf-8"))
                   if combination_path.is_file() else None)
    path = build(report, args.output, args.output.parent, combination)
    print(f"отчёт: {path}")


if __name__ == "__main__":
    main()
