"""Patch E.1: методическая UI-модель, evidence, аудит и экспорт."""
from __future__ import annotations

import inspect
import os
import sqlite3

from openpyxl import load_workbook

from protocol import comparison as cmp
from protocol import comparison_view as view
from protocol import db as protocol_db
from protocol import feature_model as model
from protocol.comparison_export import export_comparison_xlsx
from tests.method_feature_helpers import qualified_feature

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")


def _case(tmp_path):
    pdb = protocol_db.ProtocolDB(str(tmp_path / "case.db"))
    pid = pdb.create_project("Дело E.1", expert_name="Эксперт 1")
    a = pdb.add_document(pid, "TEXT_A.txt", protocol_db.ROLE_DISPUTED, "a")
    b = pdb.add_document(pid, "TEXT_B.txt", protocol_db.ROLE_SAMPLE, "b")
    pdb.save_layer(a, protocol_db.LAYER_ORIGINAL,
                   "Первое предложение. фраг Тема находится здесь.")
    pdb.save_layer(b, protocol_db.LAYER_ORIGINAL,
                   "Иной абзац. фраг Тема тоже находится здесь.")
    return pdb, pid, a, b


def _add_aux(pdb, document_id, value="0.55", fragment=None):
    pdb.save_feature_candidates(document_id, [{
        "group_name": "языковые", "subgroup": "лексические",
        "kind": "счётчик", "label": "Лексическое разнообразие (TTR)",
        "value": value, "fragment": fragment, "source": "statistics",
        "role": model.AUX_METRIC, "source_kind": model.SOURCE_ENGINEERING,
        "id_value": "", "reliability": "",
    }])


def test_01_explicit_method_high_is_preserved():
    assert view._source_informativeness(  # noqa: SLF001
        {"reference_informativeness": "высокая"}) == view.INFO_HIGH


def test_02_missing_method_informativeness_is_not_specified():
    assert view._source_informativeness({}) == view.INFO_NOT_SPECIFIED  # noqa: SLF001


def test_03_aux_does_not_enter_method_count(tmp_path):
    pdb, pid, a, b = _case(tmp_path)
    _add_aux(pdb, a); _add_aux(pdb, b)
    rows = view.build_feature_views(pdb, pid, a, b)
    summary = view.summarize(rows)
    assert summary["aux"] == 1
    assert sum(summary["method"].values()) == 0


def test_04_05_expert_fields_are_empty_before_review(tmp_path):
    pdb, pid, a, b = _case(tmp_path)
    qualified_feature(pdb, pid, a, "Тема")
    qualified_feature(pdb, pid, b, "Тема")
    cmp.auto_match(pdb, pid, a, b)
    row = next(r for r in view.build_feature_views(pdb, pid, a, b)
               if r.role == view.METHOD_FEATURE)
    assert row.expert_identification_value == view.VALUE_NOT_SET
    assert row.expert_status == view.NOT_REVIEWED


def test_06_evidence_span_maps_to_correct_document_and_context(tmp_path):
    pdb, pid, a, b = _case(tmp_path)
    qualified_feature(pdb, pid, a, "Тема")
    qualified_feature(pdb, pid, b, "Тема")
    cmp.auto_match(pdb, pid, a, b)
    row = next(r for r in view.build_feature_views(pdb, pid, a, b)
               if r.role == view.METHOD_FEATURE)
    assert row.evidence_spans_text1[0].document_id == a
    assert row.evidence_spans_text2[0].document_id == b
    assert "находится здесь" in row.evidence_spans_text1[0].context


def test_07_aggregate_aux_has_no_fake_highlight(tmp_path):
    pdb, pid, a, b = _case(tmp_path)
    _add_aux(pdb, a); _add_aux(pdb, b)
    row = next(r for r in view.build_feature_views(pdb, pid, a, b)
               if r.role == view.AUX_METRIC)
    assert row.evidence_spans_text1 == []
    assert row.evidence_spans_text2 == []
    assert row.has_local_evidence is False


def test_08_role_filters_separate_method_and_aux(tmp_path):
    pdb, pid, a, b = _case(tmp_path)
    qualified_feature(pdb, pid, a, "Тема"); qualified_feature(pdb, pid, b, "Тема")
    _add_aux(pdb, a); _add_aux(pdb, b); cmp.auto_match(pdb, pid, a, b)
    rows = view.build_feature_views(pdb, pid, a, b)
    assert all(r.role == view.METHOD_FEATURE for r in
               view.filter_views(rows, role=view.METHOD_FEATURE))
    assert all(r.role == view.AUX_METRIC for r in
               view.filter_views(rows, role=view.AUX_METRIC))


def test_09_decision_audit_has_old_new_values_and_versions(tmp_path):
    pdb, pid, a, b = _case(tmp_path)
    qualified_feature(pdb, pid, a, "Тема"); qualified_feature(pdb, pid, b, "Тема")
    cmp.auto_match(pdb, pid, a, b); position = pdb.fetch_comparisons(a, b)[0]
    cmp.reject(pdb, pid, a, b, position["position_key"], "средняя", "контекст ложный",
               program_version="5.0", expert_id="Эксперт 1")
    event = pdb.fetch_comparison_decisions(a, b)[0]
    assert event["expert_id"] == "Эксперт 1"
    assert event["old_status"] == cmp.STATUS_AUTO
    assert event["new_status"] == cmp.STATUS_REJECTED
    assert event["old_identification_value"] == ""
    assert event["new_identification_value"] == "средняя"
    assert event["comment"] == "контекст ложный"
    assert event["program_version"] == "5.0"
    assert event["method_registry_version"].startswith("sha256:")


def test_10_export_keeps_aux_on_separate_sheet(tmp_path):
    pdb, pid, a, b = _case(tmp_path)
    qualified_feature(pdb, pid, a, "Тема"); qualified_feature(pdb, pid, b, "Тема")
    _add_aux(pdb, a); _add_aux(pdb, b); cmp.auto_match(pdb, pid, a, b)
    position = pdb.fetch_comparisons(a, b)[0]
    cmp.decide(pdb, pid, a, b, position["position_key"],
               match_type=cmp.MATCH_COINCIDENCE, level="НН")
    output = tmp_path / "comparison.xlsx"
    export_comparison_xlsx(pdb, pid, a, b, str(output))
    workbook = load_workbook(output, read_only=True)
    assert workbook.sheetnames == ["METHOD — принято", "AUX"]
    assert workbook["METHOD — принято"]["B2"].value == model.METHOD_FEATURE
    assert workbook["AUX"]["B2"].value == model.AUX_METRIC


def test_11_ui_has_no_mechanical_insufficiency_formula():
    from ui.tabs import comparative_research_tab

    source = inspect.getsource(comparative_research_tab)
    assert "из требуемых" not in source
    assert "до ориентира" not in source
    assert "N < 20" not in source


def test_evidence_pane_navigation_is_clickable():
    from PyQt6.QtWidgets import QApplication
    from ui.tabs.comparative_research_tab import EvidenceTextPane

    app = QApplication.instance() or QApplication([])
    pane = EvidenceTextPane("Текст")
    pane.editor.setPlainText("один два один")
    spans = view.locate_evidence_spans(7, "один два один", ["один"])
    pane.set_spans(spans)
    assert pane.counter.text() == "1 из 2"
    assert len(pane.editor.extraSelections()) == 2
    pane.next_btn.click(); app.processEvents()
    assert pane.counter.text() == "2 из 2"


def test_legacy_high_information_is_preserved_but_not_promoted(tmp_path):
    pdb, pid, a, b = _case(tmp_path)
    with sqlite3.connect(pdb.path) as conn:
        conn.execute("ALTER TABLE comparisons ADD COLUMN high_information INTEGER")
        conn.execute(
            "INSERT INTO comparisons(project_id,pair_doc_a,pair_doc_b,position_key,"
            "label,match_type,status,created_at,high_information) VALUES(?,?,?,?,?,?,?,?,?)",
            (pid, a, b, "legacy", "Legacy", cmp.MATCH_COINCIDENCE,
             cmp.STATUS_AUTO, "2026", 1))
    migrated = protocol_db.ProtocolDB(pdb.path)
    row = migrated.fetch_comparisons(a, b)[0]
    assert row["legacy_high_information"] == 1
    rendered = view.build_feature_views(migrated, pid, a, b)[0]
    assert rendered.method_informativeness == view.INFO_NOT_SPECIFIED
