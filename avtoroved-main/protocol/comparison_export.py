"""DOCX/XLSX-приложения сравнительного исследования."""
from __future__ import annotations

from pathlib import Path

from protocol import comparison_view as view
from protocol.ingest import file_sha256


HEADERS = (
    "Название", "Роль", "Группа", "Уровень", "Текст 1", "Текст 2",
    "Отношение", "Методический источник", "Информативность по методике",
    "Идентификационная значимость", "Комментарий", "Evidence",
)


def _evidence_examples(row: view.ComparisonFeatureView) -> str:
    values = [f"Текст 1: {span.context}" for span in row.evidence_spans_text1]
    values += [f"Текст 2: {span.context}" for span in row.evidence_spans_text2]
    return " | ".join(dict.fromkeys(values))


def _source(row: view.ComparisonFeatureView) -> str:
    trace = row.method_trace
    parts = [trace.title, trace.section, trace.subsection, trace.page,
             trace.quote_or_summary]
    return "; ".join(part for part in parts if part)


def comparison_export_rows(pdb, project_id: int, doc_a: int, doc_b: int) -> dict:
    """Развести принятый METHOD-комплекс и AUX без пере-классификации."""
    rows = view.build_feature_views(pdb, project_id, doc_a, doc_b)

    def serialize(row):
        return (
            row.feature_name, row.role, f"{row.method_group} / {row.method_subgroup}",
            row.individualization_level, row.value_text1, row.value_text2,
            row.comparison_relation, _source(row), row.method_informativeness,
            row.expert_identification_value, row.notes, _evidence_examples(row),
        )

    return {
        "method": [serialize(row) for row in rows
                   if row.role == view.METHOD_FEATURE
                   and row.expert_status == view.ACCEPTED],
        "aux": [serialize(row) for row in rows if row.role == view.AUX_METRIC],
    }


def export_comparison_docx(pdb, project_id: int, doc_a: int, doc_b: int,
                           filepath: str, program_version: str | None = None) -> dict:
    from docx import Document
    from docx.shared import Pt

    payload = comparison_export_rows(pdb, project_id, doc_a, doc_b)
    document = Document(); document.styles["Normal"].font.name = "Times New Roman"
    document.styles["Normal"].font.size = Pt(10)
    document.add_heading("Сравнительное исследование", level=1)
    document.add_paragraph(
        "В основной раздел включены только принятые экспертом METHOD_FEATURE. "
        "Вспомогательные показатели приведены отдельно и не входят в методический комплекс.")
    _add_docx_table(document, "Принятые методические признаки", payload["method"])
    _add_docx_table(document, "Вспомогательные объективизирующие показатели",
                    payload["aux"])
    document.save(filepath)
    return _record_export(pdb, project_id, doc_a, doc_b, filepath,
                          "DOCX", program_version)


def _add_docx_table(document, title: str, rows: list[tuple]) -> None:
    document.add_heading(title, level=2)
    if not rows:
        document.add_paragraph("Нет данных."); return
    table = document.add_table(rows=1, cols=len(HEADERS)); table.style = "Table Grid"
    for index, header in enumerate(HEADERS): table.rows[0].cells[index].text = header
    for values in rows:
        cells = table.add_row().cells
        for index, value in enumerate(values): cells[index].text = str(value or "—")


def export_comparison_xlsx(pdb, project_id: int, doc_a: int, doc_b: int,
                           filepath: str, program_version: str | None = None) -> dict:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    payload = comparison_export_rows(pdb, project_id, doc_a, doc_b)
    workbook = Workbook(); workbook.remove(workbook.active)
    for title, key in (("METHOD — принято", "method"), ("AUX", "aux")):
        sheet = workbook.create_sheet(title)
        sheet.append(HEADERS)
        for cell in sheet[1]: cell.font = Font(bold=True)
        for row in payload[key]: sheet.append(row)
        sheet.freeze_panes = "A2"; sheet.auto_filter.ref = sheet.dimensions
        for column in sheet.columns:
            width = min(55, max(12, max(len(str(cell.value or "")) for cell in column) + 2))
            sheet.column_dimensions[column[0].column_letter].width = width
    workbook.save(filepath)
    return _record_export(pdb, project_id, doc_a, doc_b, filepath,
                          "XLSX", program_version)


def _record_export(pdb, project_id, doc_a, doc_b, filepath, kind, program_version):
    sha = file_sha256(filepath)
    pdb.record_report(project_id, str(Path(filepath)), sha, pair_doc_a=doc_a,
                      pair_doc_b=doc_b, program_version=program_version)
    pdb.log_action(
        "экспортировано сравнительное исследование", project_id=project_id,
        details={"pair_doc_a": doc_a, "pair_doc_b": doc_b, "формат": kind,
                 "файл": str(Path(filepath)), "sha256": sha},
        program_version=program_version)
    return {"filepath": str(Path(filepath)), "sha256": sha, "format": kind}
